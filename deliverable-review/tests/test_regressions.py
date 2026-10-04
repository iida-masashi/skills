"""Regression tests for bugs found in the 2026-10 review.

Each test builds a minimal document in tmp_path; no network / external API.
"""
from __future__ import annotations

import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

SKILL_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS_DIR = SKILL_ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

docx = pytest.importorskip("docx")
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

import checkers  # noqa: E402
import extractors  # noqa: E402
import internal_content  # noqa: E402
import llm_review  # noqa: E402
import metadata  # noqa: E402
import numeric_integrity as N  # noqa: E402
import patterns as P  # noqa: E402
import ai_check_extract  # noqa: E402


# ------------------------------------------------------------
# Sanitize
# ------------------------------------------------------------

def _docx_with_comment(path: Path) -> None:
    d = docx.Document()
    d.add_paragraph("本文その1")
    p = d.add_paragraph("本文その2")
    d.add_comment(p.runs, text="TODO 先方には言わない", author="Reviewer")
    d.save(path)


def test_docx_sanitize_with_comments_stays_openable(tmp_path):
    src, dst = tmp_path / "c.docx", tmp_path / "c_san.docx"
    _docx_with_comment(src)
    actions = metadata.sanitize(str(src), str(dst), ".docx")

    with zipfile.ZipFile(dst) as z:
        names = z.namelist()
        rels = z.read("word/_rels/document.xml.rels").decode()
        ct = z.read("[Content_Types].xml").decode()
    assert "word/comments.xml" not in names
    assert "comments" not in rels
    assert "comments" not in ct
    d = docx.Document(dst)  # must not raise
    assert [p.text for p in d.paragraphs] == ["本文その1", "本文その2"]
    assert actions[-1].startswith("verify: OK")


_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def test_accept_tracked_changes_keeps_text_around_self_closing_del():
    xml = (
        f'<w:document {_W}><w:body>'
        '<w:p><w:pPr><w:rPr><w:del w:id="1" w:author="a"/></w:rPr></w:pPr>'
        '<w:r><w:t>KEEP-1</w:t></w:r></w:p>'
        '<w:p><w:r><w:t>KEEP-2</w:t></w:r>'
        '<w:del w:id="2"><w:r><w:delText>GONE</w:delText></w:r></w:del>'
        '<w:ins w:id="3"><w:r><w:t>ADDED</w:t></w:r></w:ins></w:p>'
        '</w:body></w:document>'
    ).encode()
    out, n_ins, n_del = metadata._accept_tracked_changes(xml)
    text = out.decode()
    assert "KEEP-1" in text and "KEEP-2" in text and "ADDED" in text
    assert "GONE" not in text
    assert "<w:del" not in text and "<w:ins" not in text
    assert (n_ins, n_del) == (1, 2)


def test_pptx_sanitize_clears_company_and_custom_props_keeps_msip(tmp_path):
    src, dst = tmp_path / "m.pptx", tmp_path / "m_san.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(src)
    # inject Company + custom.xml (one ordinary prop + one sensitivity label)
    custom = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="ClientCode">'
        '<vt:lpwstr>ACME-001</vt:lpwstr></property>'
        '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="3" name="MSIP_Label_x_Enabled">'
        '<vt:lpwstr>true</vt:lpwstr></property></Properties>'
    )
    tmp = tmp_path / "tmp.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = data.replace(b"</Properties>", b"<Company>Fortience</Company></Properties>")
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", (
                    b'<Override PartName="/docProps/custom.xml" ContentType='
                    b'"application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>'))
            if item.filename == "_rels/.rels":
                data = data.replace(b"</Relationships>", (
                    b'<Relationship Id="rIdC" Type="http://schemas.openxmlformats.org/officeDocument/'
                    b'2006/relationships/custom-properties" Target="docProps/custom.xml"/></Relationships>'))
            zout.writestr(item, data)
        zout.writestr("docProps/custom.xml", custom)

    found = {f.category for f in metadata.check_metadata(extractors.extract(str(tmp)))}
    assert {"app-company", "custom-properties"} <= found

    metadata.sanitize(str(tmp), str(dst), ".pptx")
    with zipfile.ZipFile(dst) as z:
        assert b"Fortience" not in z.read("docProps/app.xml")
        cx = z.read("docProps/custom.xml")
    assert b"ClientCode" not in cx and b"MSIP_Label_x_Enabled" in cx
    Presentation(dst)


# ------------------------------------------------------------
# Hyperlinks / citations / URL parsing
# ------------------------------------------------------------

def test_hyperlink_target_utm_detected(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = "出典"
    r.hyperlink.address = "https://example.com/?utm_source=chatgpt.com"
    p = tmp_path / "h.pptx"
    prs.save(p)
    cats = [f.category for f in checkers.check_url_contamination(extractors.extract(str(p)))]
    assert cats == ["AI-origin-query-param"]


def test_copyright_footer_is_not_a_citation(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = "当社シェアは35%で国内第1位"
    s.shapes.add_textbox(Inches(1), Inches(6), Inches(5), Inches(1)).text_frame.text = "© 2026 Example Inc."
    p = tmp_path / "c.pptx"
    prs.save(p)
    claims = checkers.check_verifiable_claims(extractors.extract(str(p)))
    assert any("35%" in f.evidence for f in claims)


def test_fullwidth_paren_not_part_of_url():
    assert P.extract_all_urls("（出典: https://www.meti.go.jp/report）") == ["https://www.meti.go.jp/report"]


def test_business_date_is_not_cutoff_mention():
    assert P.find_ai_traces("2024年3月時点の売上") == []
    assert [c for c, _ in P.find_ai_traces("2023年4月時点の私の知識では")] == ["knowledge-cutoff-mention"]


def test_docx_citation_scoped_to_section(tmp_path):
    d = docx.Document()
    d.add_heading("市場", level=1)
    d.add_paragraph("市場規模は1,200億円（出典: 経産省）")
    d.add_heading("当社", level=1)
    d.add_paragraph("当社シェアは国内第1位で、売上高は500億円。")
    p = tmp_path / "s.docx"
    d.save(p)
    doc = extractors.extract(str(p))
    assert len(doc.location_flags) == 2
    claims = checkers.check_verifiable_claims(doc)
    assert any("第1位" in f.evidence for f in claims)
    assert not any("1,200億円" in f.evidence for f in claims)


# ------------------------------------------------------------
# False positives
# ------------------------------------------------------------

def _evidence(rows):
    return [f.evidence for f in N._check_table_totals(rows, "T", 1, "t")]


def test_table_year_header_not_summed():
    rows = [["項目", "2023年", "2024年"], ["売上", "100", "120"], ["費用", "60", "70"], ["合計", "160", "190"]]
    assert _evidence(rows) == []


def test_table_keikaku_row_not_total():
    rows = [["項目", "金額"], ["調査", "20"], ["計画策定", "30"], ["実行", "50"]]
    assert _evidence(rows) == []


def test_table_subtotals_not_double_counted():
    rows = [["項目", "金額"], ["A", "10"], ["B", "20"], ["小計", "30"], ["C", "40"], ["合計", "70"]]
    assert _evidence(rows) == []
    bad = [["項目", "金額"], ["A", "10"], ["B", "20"], ["小計", "35"], ["合計", "35"]]
    assert any("labeled=35.0" in e for e in _evidence(bad))


def test_table_triangle_negative():
    rows = [["項目", "金額"], ["A", "▲20"], ["B", "50"], ["合計", "30"]]
    assert _evidence(rows) == []


def test_table_ratio_column_excluded_from_row_total():
    rows = [["項目", "国内", "海外", "構成比", "合計"], ["売上", "60", "40", "50%", "100"]]
    assert _evidence(rows) == []


def test_danger_words_graded():
    sev, _ = internal_content._grade("前回は仮説ベースで説明", "INFO")
    assert sev == "INFO"
    sev, _ = internal_content._grade("原価の話は避けたい", "INFO")
    assert sev == "MEDIUM"
    sev, _ = internal_content._grade("ここは社内限りで", "INFO")
    assert sev == "HIGH"


def test_hidden_slide_reported_once(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "メモ"
    s._element.set("show", "0")
    p = tmp_path / "hid.pptx"
    prs.save(p)
    fs = checkers.run_all(extractors.extract(str(p)), skip_liveness=True)
    assert [f.checker for f in fs if f.category == "hidden-slide"] == ["internal-content"]


# ------------------------------------------------------------
# CLI / LLM payload
# ------------------------------------------------------------

def test_cli_runs_strategy_and_fail_on(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "現状"
    prs.core_properties.author = "Someone"
    p = tmp_path / "x.pptx"
    prs.save(p)
    fs = checkers.run_all(extractors.extract(str(p)), skip_liveness=True)
    assert any(f.checker == "strategy" for f in fs)

    base = [sys.executable, str(SCRIPTS_DIR / "review.py"), str(p), "--skip-liveness",
            "--out-dir", str(tmp_path)]
    assert subprocess.run(base, capture_output=True, timeout=60).returncode == 0
    assert subprocess.run(base + ["--fail-on", "HIGH"], capture_output=True, timeout=60).returncode == 1
    report = (tmp_path / "x_review.md").read_text(encoding="utf-8")
    assert "戦略コンサル品質 (strategy)" in report


def test_llm_prompt_includes_docx_body(tmp_path):
    d = docx.Document()
    d.add_paragraph("市場規模は1,200億円と推計される")
    p = tmp_path / "l.docx"
    d.save(p)
    j = tmp_path / "l.json"
    ai_check_extract.write_ai_check_json(str(p), str(j))
    import json
    prompt = llm_review._build_user_prompt(json.loads(j.read_text(encoding="utf-8")))
    assert "市場規模は1,200億円" in prompt


def test_pdf_sanitize_removes_info(tmp_path):
    pypdf = pytest.importorskip("pypdf")
    w = pypdf.PdfWriter()
    w.add_blank_page(200, 200)
    w.add_metadata({"/Author": "Someone", "/Company": "X"})
    src = tmp_path / "m.pdf"
    w.write(src)
    actions = metadata.sanitize(str(src), str(tmp_path / "m_san.pdf"), ".pdf")
    assert actions[-1].startswith("verify: OK")


def test_pptx_comment_mapped_by_slide_rels_and_sanitized(tmp_path):
    prs = Presentation()
    for _ in range(2):
        prs.slides.add_slide(prs.slide_layouts[6])
    src = tmp_path / "cm.pptx"
    prs.save(src)
    cm = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
          '<p:cm authorId="0" dt="2026-10-01T00:00:00" idx="1"><p:pos x="10" y="10"/>'
          '<p:text>社内限りの数字です</p:text></p:cm></p:cmLst>')
    au = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
          '<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
          '<p:cmAuthor id="0" name="Reviewer" initials="R" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>')
    rel_c = (b'<Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/'
             b'2006/relationships/comments" Target="../comments/comment1.xml"/></Relationships>')
    rel_a = (b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/'
             b'2006/relationships/commentAuthors" Target="commentAuthors.xml"/></Relationships>')
    ct = (b'<Override PartName="/ppt/comments/comment1.xml" ContentType="application/'
          b'vnd.openxmlformats-officedocument.presentationml.comments+xml"/>'
          b'<Override PartName="/ppt/commentAuthors.xml" ContentType="application/'
          b'vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"/></Types>')
    tmp = tmp_path / "cm2.pptx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "ppt/slides/_rels/slide2.xml.rels":
                data = data.replace(b"</Relationships>", rel_c)
            if it.filename == "ppt/_rels/presentation.xml.rels":
                data = data.replace(b"</Relationships>", rel_a)
            if it.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", ct)
            zout.writestr(it, data)
        zout.writestr("ppt/comments/comment1.xml", cm)
        zout.writestr("ppt/commentAuthors.xml", au)

    fs = internal_content.run_internal_content(extractors.extract(str(tmp)))
    # comment1.xml belongs to slide 2 (not slide 1)
    assert [(f.location_label, f.severity) for f in fs] == [("Slide 2", "HIGH")]
    actions = metadata.sanitize(str(tmp), str(tmp_path / "cm_san.pptx"), ".pptx")
    assert actions[-1].startswith("verify: OK")
